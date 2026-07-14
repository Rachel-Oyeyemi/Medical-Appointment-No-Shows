from __future__ import annotations

from pathlib import Path
import json
import textwrap
import nbformat as nbf

ROOT = Path(__file__).resolve().parent


def write(path: str, content: str) -> None:
    target = ROOT / path
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(textwrap.dedent(content).lstrip(), encoding="utf-8")


write("requirements.txt", """
pandas>=2.2,<3
numpy>=2.0,<3
scikit-learn>=1.5,<2
xgboost>=2.1,<4
joblib>=1.4,<2
matplotlib>=3.9,<4
streamlit>=1.40,<2
kagglehub>=0.3,<2
kaggle>=1.6,<2
python-pptx>=1.0,<2
pytest>=8,<10
nbformat>=5.10,<6
""")
write(".gitignore", """
.venv/
__pycache__/
*.py[cod]
.pytest_cache/
.ipynb_checkpoints/
.env
kaggle.json
.DS_Store
.streamlit/secrets.toml
data/raw/*.csv
data/processed/*.csv
visuals/*.png
""")
write("LICENSE", """
MIT License

Copyright (c) 2026 Rachel Oyeyemi

Permission is hereby granted, free of charge, to any person obtaining a copy of this software and associated documentation files (the "Software"), to deal in the Software without restriction, including without limitation the rights to use, copy, modify, merge, publish, distribute, sublicense, and/or sell copies of the Software.

THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR IMPLIED.
""")
write("src/__init__.py", '"""Medical Appointment No Shows project package."""\n')
write("src/utils.py", '''
"""Shared project utilities."""
from __future__ import annotations
import json, logging
from pathlib import Path
from typing import Any

def configure_logging() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(name)s | %(message)s")

def save_json(payload: dict[str, Any], path: str | Path) -> None:
    output = Path(path); output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")

def load_json(path: str | Path) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))
''')
write("src/generate_demo_data.py", '''
"""Generate deterministic non-identifying appointment data."""
from __future__ import annotations
import argparse
from pathlib import Path
import numpy as np
import pandas as pd

NEIGHBOURHOODS = ["JARDIM CAMBURI","MARIA ORTIZ","RESISTÊNCIA","JARDIM DA PENHA","ITARARÉ","CENTRO","TABUAZEIRO","SANTA MARTHA","BONFIM","JESUS DE NAZARETH","SANTO ANTÔNIO","SÃO PEDRO","CARATOÍRA","ANDORINHAS","NOVA PALESTINA","DA PENHA","ROMÃO","GOIABEIRAS","ILHA DO PRÍNCIPE","MARUÍPE","BELA VISTA","FORTE SÃO JOÃO","SÃO CRISTÓVÃO","CONQUISTA","REDENÇÃO"]

def generate_demo(n_rows: int = 30000, seed: int = 42) -> pd.DataFrame:
    rng=np.random.default_rng(seed)
    patient_id=rng.choice(np.arange(10_000_000,10_012_000),n_rows)
    appointment_id=np.arange(5_000_000,5_000_000+n_rows)
    appointment_day=pd.Timestamp("2016-04-29")+pd.to_timedelta(rng.integers(0,42,n_rows),unit="D")
    waiting=np.clip(np.where(rng.random(n_rows)<.42,0,rng.gamma(2.2,8,n_rows).astype(int)),0,90)
    hour=np.clip(np.round(rng.normal(10.5,3,n_rows)).astype(int),6,19)
    scheduled=appointment_day-pd.to_timedelta(waiting,unit="D")+pd.to_timedelta(hour,unit="h")+pd.to_timedelta(rng.integers(0,60,n_rows),unit="m")
    gender=rng.choice(["F","M"],p=[.65,.35],size=n_rows)
    age=np.clip(np.round(rng.gamma(2.3,17,n_rows)).astype(int),0,102)
    hood=rng.choice(NEIGHBOURHOODS,n_rows)
    scholarship=rng.binomial(1,np.clip(.06+.07*((age>=18)&(age<=50)),0,1))
    hypertension=rng.binomial(1,1/(1+np.exp(-(-6.2+.09*age))))
    diabetes=rng.binomial(1,1/(1+np.exp(-(-7.2+.085*age))))
    alcoholism=rng.binomial(1,np.clip(.015+.04*((age>=25)&(age<=65)),0,.15))
    handicap=rng.choice([0,1,2],p=[.978,.020,.002],size=n_rows)
    sms=rng.binomial(1,np.clip(.06+.018*waiting+.10*(waiting>=7),.04,.82))
    effects={name:value for name,value in zip(NEIGHBOURHOODS,rng.normal(0,.38,len(NEIGHBOURHOODS)))}
    hood_effect=np.array([effects[x] for x in hood]); seg_a=np.isin(hood,NEIGHBOURHOODS[:6]); seg_b=np.isin(hood,NEIGHBOURHOODS[6:12])
    dow=appointment_day.dayofweek.to_numpy(); young=(age>=15)&(age<=35); older=age>=60
    logit=(-2.20+.030*waiting+.55*(waiting>=21)+.35*(waiting>=45)+.38*young-.35*older+.23*scholarship+.43*alcoholism+.20*(hour>=16)+.18*(dow==0)+hood_effect-.52*(waiting==0)+.55*((waiting>=14)&(sms==0))-.30*((waiting>=7)&(sms==1))+.30*((age<10)&(waiting>=30))+.95*(seg_a&young)-.65*(seg_a&older)+.75*(seg_b&(waiting>=15))+.60*(((dow==0)|(dow==4))&(waiting>=8)&(waiting<=30))+.55*((hour>=15)&(waiting>=21))-.18*hypertension-.12*diabetes+rng.normal(0,.72,n_rows))
    no_show=np.where(rng.random(n_rows)<1/(1+np.exp(-logit)),"Yes","No")
    return pd.DataFrame({"PatientId":patient_id,"AppointmentID":appointment_id,"Gender":gender,"ScheduledDay":scheduled.strftime("%Y-%m-%dT%H:%M:%SZ"),"AppointmentDay":appointment_day.strftime("%Y-%m-%dT00:00:00Z"),"Age":age,"Neighbourhood":hood,"Scholarship":scholarship,"Hipertension":hypertension,"Diabetes":diabetes,"Alcoholism":alcoholism,"Handcap":handicap,"SMS_received":sms,"No-show":no_show})

def main():
    p=argparse.ArgumentParser(); p.add_argument("--rows",type=int,default=30000); p.add_argument("--output",default="data/sample_data/demo_appointments.csv"); a=p.parse_args()
    out=Path(a.output); out.parent.mkdir(parents=True,exist_ok=True); generate_demo(a.rows).to_csv(out,index=False)
    print(f"Saved {a.rows:,} synthetic appointments to {out}")
if __name__=="__main__": main()
''')
write("src/download_data.py", '''
"""Download the official dataset from Kaggle."""
from __future__ import annotations
import argparse, logging, shutil, subprocess
from pathlib import Path
from utils import configure_logging
LOGGER=logging.getLogger(__name__); HANDLE="joniarroba/noshowappointments"

def kagglehub_download(output: Path) -> Path:
    import kagglehub
    cache=Path(kagglehub.dataset_download(HANDLE)); files=list(cache.rglob("*.csv"))
    if not files: raise FileNotFoundError("No CSV in KaggleHub download")
    output.mkdir(parents=True,exist_ok=True); target=output/"appointments.csv"; shutil.copy2(files[0],target); return target

def cli_download(output: Path) -> Path:
    output.mkdir(parents=True,exist_ok=True); subprocess.run(["kaggle","datasets","download","-d",HANDLE,"-p",str(output),"--unzip"],check=True)
    files=list(output.glob("*.csv"));
    if not files: raise FileNotFoundError("No CSV produced by Kaggle CLI")
    target=output/"appointments.csv"; files[0].replace(target) if files[0]!=target else None; return target

def main():
    p=argparse.ArgumentParser(); p.add_argument("--output-dir",default="data/raw"); p.add_argument("--source",choices=["auto","kagglehub","cli"],default="auto"); a=p.parse_args(); configure_logging(); errors=[]
    methods=[kagglehub_download,cli_download] if a.source=="auto" else [kagglehub_download if a.source=="kagglehub" else cli_download]
    for method in methods:
        try: LOGGER.info("Downloaded to %s",method(Path(a.output_dir))); return
        except Exception as exc: errors.append(f"{method.__name__}: {exc}")
    raise RuntimeError("Authenticate to Kaggle before downloading. "+" | ".join(errors))
if __name__=="__main__": main()
''')
write("src/preprocess.py", '''
"""Validate and clean raw medical appointment records."""
from __future__ import annotations
import argparse, logging
from pathlib import Path
import pandas as pd
from utils import configure_logging, save_json
LOGGER=logging.getLogger(__name__)
REQUIRED={"PatientId","AppointmentID","Gender","ScheduledDay","AppointmentDay","Age","Neighbourhood","Scholarship","Hipertension","Diabetes","Alcoholism","Handcap","SMS_received","No-show"}
RENAME={"PatientId":"patient_id","AppointmentID":"appointment_id","Gender":"gender","ScheduledDay":"scheduled_datetime","AppointmentDay":"appointment_date","Age":"age","Neighbourhood":"neighbourhood","Scholarship":"scholarship","Hipertension":"hypertension","Diabetes":"diabetes","Alcoholism":"alcoholism","Handcap":"handicap","SMS_received":"sms_received","No-show":"no_show"}

def preprocess_dataframe(frame: pd.DataFrame):
    missing=REQUIRED.difference(frame.columns)
    if missing: raise ValueError(f"Missing required columns: {sorted(missing)}")
    data=frame.copy().rename(columns=RENAME); input_rows=len(data); missing_cells=int(data.isna().sum().sum()); duplicates=int(data.duplicated().sum())
    data["scheduled_datetime"]=pd.to_datetime(data["scheduled_datetime"],utc=True,errors="coerce"); data["appointment_date"]=pd.to_datetime(data["appointment_date"],utc=True,errors="coerce").dt.normalize()
    data["gender"]=data.gender.astype("string").str.upper().str.strip(); data["neighbourhood"]=data.neighbourhood.astype("string").str.strip().str.upper(); data["no_show"]=data.no_show.astype("string").str.lower().str.strip().map({"yes":1,"no":0})
    for col in ["patient_id","appointment_id","age","scholarship","hypertension","diabetes","alcoholism","handicap","sms_received"]: data[col]=pd.to_numeric(data[col],errors="coerce")
    data=data.drop_duplicates(subset=["appointment_id"]).dropna(subset=["scheduled_datetime","appointment_date","age","gender","neighbourhood","no_show"]); data=data[data.gender.isin(["F","M"])]
    bad_age=int(((data.age<0)|(data.age>110)).sum()); data=data[data.age.between(0,110)]
    data["waiting_days"]=(data.appointment_date.dt.tz_localize(None)-data.scheduled_datetime.dt.tz_convert(None).dt.normalize()).dt.days; bad_wait=int((data.waiting_days<0).sum()); data=data[data.waiting_days>=0]
    for col in ["scholarship","hypertension","diabetes","alcoholism","sms_received"]: data[col]=data[col].fillna(0).clip(0,1).astype("int8")
    data["handicap"]=data.handicap.fillna(0).clip(lower=0).astype("int8"); data["age"]=data.age.astype("int16"); data["patient_id"]=data.patient_id.astype("int64"); data["appointment_id"]=data.appointment_id.astype("int64"); data["no_show"]=data.no_show.astype("int8")
    data=data.sort_values(["appointment_date","scheduled_datetime","appointment_id"]).reset_index(drop=True)
    quality={"input_rows":input_rows,"output_rows":len(data),"rows_removed":input_rows-len(data),"input_columns":frame.shape[1],"missing_cells_before":missing_cells,"exact_duplicate_rows_before":duplicates,"invalid_age_rows_removed":bad_age,"negative_wait_rows_removed":bad_wait,"target_positive_rows":int(data.no_show.sum()),"target_positive_rate":float(data.no_show.mean()),"date_min":str(data.appointment_date.min().date()),"date_max":str(data.appointment_date.max().date()),"unique_patients":int(data.patient_id.nunique()),"unique_neighbourhoods":int(data.neighbourhood.nunique())}
    return data,quality

def main():
    p=argparse.ArgumentParser(); p.add_argument("--input",default="data/raw/appointments.csv"); p.add_argument("--output",default="data/processed/clean_appointments.csv"); p.add_argument("--quality-output",default="reports/data_quality.json"); a=p.parse_args(); configure_logging()
    clean,quality=preprocess_dataframe(pd.read_csv(a.input,low_memory=False)); out=Path(a.output); out.parent.mkdir(parents=True,exist_ok=True); clean.to_csv(out,index=False); save_json(quality,a.quality_output); LOGGER.info("Saved %s rows",len(clean))
if __name__=="__main__": main()
''')
write("src/feature_engineering.py", '''
"""Create leakage-aware features available before the appointment."""
from __future__ import annotations
import argparse
from pathlib import Path
import numpy as np, pandas as pd
TARGET="no_show"; NUMERIC_FEATURES=["age","waiting_days","schedule_hour","chronic_count","handicap"]; BINARY_FEATURES=["scholarship","hypertension","diabetes","alcoholism","sms_received","same_day"]; CATEGORICAL_FEATURES=["gender","neighbourhood","appointment_dow","schedule_dow","appointment_month","age_group","waiting_bucket"]; MODEL_FEATURES=NUMERIC_FEATURES+BINARY_FEATURES+CATEGORICAL_FEATURES

def engineer_features(frame: pd.DataFrame) -> pd.DataFrame:
    data=frame.copy(); data["scheduled_datetime"]=pd.to_datetime(data.scheduled_datetime,utc=True); data["appointment_date"]=pd.to_datetime(data.appointment_date,utc=True)
    if "waiting_days" not in data: data["waiting_days"]=(data.appointment_date.dt.tz_localize(None)-data.scheduled_datetime.dt.tz_convert(None).dt.normalize()).dt.days
    data["schedule_hour"]=data.scheduled_datetime.dt.hour.astype("int8"); data["schedule_dow"]=data.scheduled_datetime.dt.day_name().str[:3]; data["appointment_dow"]=data.appointment_date.dt.day_name().str[:3]; data["appointment_month"]=data.appointment_date.dt.month.astype(str); data["same_day"]=(data.waiting_days==0).astype("int8"); data["chronic_count"]=data[["hypertension","diabetes","alcoholism"]].sum(axis=1).astype("int8")
    data["age_group"]=pd.cut(data.age,[-1,5,12,17,29,44,59,74,110],labels=["0-5","6-12","13-17","18-29","30-44","45-59","60-74","75+"]).astype("string")
    data["waiting_bucket"]=pd.cut(data.waiting_days,[-1,0,3,7,14,30,60,np.inf],labels=["same_day","1-3","4-7","8-14","15-30","31-60","61+"]).astype("string")
    for col in CATEGORICAL_FEATURES: data[col]=data[col].fillna("Unknown").astype(str)
    return data

def chronological_split(frame,validation_fraction=.15,test_fraction=.15):
    ordered=frame.sort_values("appointment_date").reset_index(drop=True); n=len(ordered); train_end=int(n*(1-validation_fraction-test_fraction)); val_end=int(n*(1-test_fraction)); return ordered.iloc[:train_end].copy(),ordered.iloc[train_end:val_end].copy(),ordered.iloc[val_end:].copy()

def main():
    p=argparse.ArgumentParser(); p.add_argument("--input",default="data/processed/clean_appointments.csv"); p.add_argument("--output",default="data/processed/model_features.csv"); a=p.parse_args(); featured=engineer_features(pd.read_csv(a.input,parse_dates=["scheduled_datetime","appointment_date"])); out=Path(a.output); out.parent.mkdir(parents=True,exist_ok=True); featured.to_csv(out,index=False)
if __name__=="__main__": main()
''')
write("src/train_model.py", '''
"""Train Logistic Regression and XGBoost no-show classifiers."""
from __future__ import annotations
import argparse, logging
from pathlib import Path
import joblib, pandas as pd
from sklearn.compose import ColumnTransformer
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from xgboost import XGBClassifier
from feature_engineering import BINARY_FEATURES,CATEGORICAL_FEATURES,MODEL_FEATURES,NUMERIC_FEATURES,TARGET,chronological_split
from utils import configure_logging,save_json
LOGGER=logging.getLogger(__name__)

def preprocessor(scale=True):
    steps=[("impute",SimpleImputer(strategy="median"))]+([("scale",StandardScaler())] if scale else [])
    return ColumnTransformer([("num",Pipeline(steps),NUMERIC_FEATURES),("bin",SimpleImputer(strategy="most_frequent"),BINARY_FEATURES),("cat",OneHotEncoder(handle_unknown="ignore",min_frequency=5),CATEGORICAL_FEATURES)])

def build_baseline(): return Pipeline([("preprocessor",preprocessor(True)),("model",LogisticRegression(max_iter=2000,class_weight="balanced",C=.8,solver="liblinear",random_state=42))])
def build_advanced(weight): return Pipeline([("preprocessor",preprocessor(False)),("model",XGBClassifier(n_estimators=500,max_depth=6,learning_rate=.04,min_child_weight=4,subsample=.85,colsample_bytree=.85,reg_alpha=.05,reg_lambda=3,gamma=.05,scale_pos_weight=weight,objective="binary:logistic",eval_metric="aucpr",tree_method="hist",random_state=42,n_jobs=4))])

def train_models(frame):
    train,validation,test=chronological_split(frame); X=train[MODEL_FEATURES]; y=train[TARGET]; positives=int(y.sum()); weight=(len(y)-positives)/positives
    return build_baseline().fit(X,y),build_advanced(weight).fit(X,y),train,validation,test,weight

def main():
    p=argparse.ArgumentParser(); p.add_argument("--input",default="data/processed/model_features.csv"); p.add_argument("--models-dir",default="models"); a=p.parse_args(); configure_logging(); frame=pd.read_csv(a.input,parse_dates=["scheduled_datetime","appointment_date"]); baseline,advanced,train,val,test,weight=train_models(frame); out=Path(a.models_dir); out.mkdir(parents=True,exist_ok=True); joblib.dump(baseline,out/"baseline_logistic_regression.joblib",compress=3); joblib.dump(advanced,out/"advanced_xgboost.joblib",compress=3); joblib.dump({"validation":val,"test":test},out/"evaluation_partitions.joblib",compress=3); save_json({"features":MODEL_FEATURES,"target":TARGET,"split_strategy":"chronological 70/15/15","scale_pos_weight":weight,"random_seed":42,"training_rows":len(train),"validation_rows":len(val),"test_rows":len(test),"identifier_policy":"PatientId and AppointmentID excluded"},out/"model_metadata.json"); LOGGER.info("Saved models")
if __name__=="__main__": main()
''')
write("src/evaluate_model.py", '''
"""Evaluate models and select transparent outreach thresholds."""
from __future__ import annotations
import argparse
from pathlib import Path
import joblib, numpy as np, pandas as pd
from sklearn.metrics import accuracy_score,average_precision_score,balanced_accuracy_score,confusion_matrix,f1_score,matthews_corrcoef,precision_score,recall_score,roc_auc_score
from feature_engineering import MODEL_FEATURES,TARGET
from utils import save_json

def choose_threshold(y,p,fn_cost=5,fp_cost=1):
    rows=[]
    for threshold in np.linspace(.05,.90,172):
        pred=(p>=threshold).astype(int); tn,fp,fn,tp=confusion_matrix(y,pred,labels=[0,1]).ravel(); rows.append({"threshold":float(threshold),"cost":float(fn_cost*fn+fp_cost*fp),"true_negatives":int(tn),"false_positives":int(fp),"false_negatives":int(fn),"true_positives":int(tp),"precision":float(precision_score(y,pred,zero_division=0)),"recall":float(recall_score(y,pred,zero_division=0)),"f1":float(f1_score(y,pred,zero_division=0))})
    best=min(rows,key=lambda x:(x["cost"],-x["recall"],-x["precision"])); return best["threshold"],rows

def evaluate(model,frame,threshold):
    y=frame[TARGET].to_numpy(); p=model.predict_proba(frame[MODEL_FEATURES])[:,1]; pred=(p>=threshold).astype(int); tn,fp,fn,tp=confusion_matrix(y,pred,labels=[0,1]).ravel(); return {"threshold":float(threshold),"accuracy":float(accuracy_score(y,pred)),"precision":float(precision_score(y,pred,zero_division=0)),"recall":float(recall_score(y,pred,zero_division=0)),"f1":float(f1_score(y,pred,zero_division=0)),"roc_auc":float(roc_auc_score(y,p)),"pr_auc":float(average_precision_score(y,p)),"balanced_accuracy":float(balanced_accuracy_score(y,pred)),"matthews_correlation":float(matthews_corrcoef(y,pred)),"specificity":float(tn/(tn+fp)),"false_positive_rate":float(fp/(tn+fp)),"predicted_outreach_rate":float(pred.mean()),"true_negatives":int(tn),"false_positives":int(fp),"false_negatives":int(fn),"true_positives":int(tp),"confusion_matrix":[[int(tn),int(fp)],[int(fn),int(tp)]]},p

def main():
    p=argparse.ArgumentParser(); p.add_argument("--models-dir",default="models"); p.add_argument("--reports-dir",default="reports"); a=p.parse_args(); md=Path(a.models_dir); rd=Path(a.reports_dir); rd.mkdir(parents=True,exist_ok=True); baseline=joblib.load(md/"baseline_logistic_regression.joblib"); advanced=joblib.load(md/"advanced_xgboost.joblib"); parts=joblib.load(md/"evaluation_partitions.joblib"); val,test=parts["validation"],parts["test"]; results=[]; thresholds={}; tables=[]
    for name,model in [("Logistic Regression",baseline),("XGBoost",advanced)]:
        threshold,rows=choose_threshold(val[TARGET],model.predict_proba(val[MODEL_FEATURES])[:,1]); thresholds[name]=threshold; tables += [{"model":name,**r} for r in rows]
        results += [{"model":name,"split":"validation",**evaluate(model,val,threshold)[0]},{"model":name,"split":"test",**evaluate(model,test,threshold)[0]}]
    recommended=max([r for r in results if r["split"]=="validation"],key=lambda r:r["pr_auc"])["model"]
    save_json({"benchmark_scope":"Deterministic synthetic appointment data; rerun on the official Kaggle CSV for final portfolio metrics.","primary_metric":"PR-AUC","models":results,"recommended_model":recommended,"threshold_cost_assumptions":{"missed_no_show":5,"unnecessary_outreach":1}},rd/"benchmark_metrics.json"); save_json(thresholds,md/"threshold_config.json"); pd.DataFrame(tables).to_csv(rd/"threshold_diagnostics.csv",index=False); print(recommended)
if __name__=="__main__": main()
''')
write("src/predict.py", '''
"""Score future appointments for supportive outreach."""
from __future__ import annotations
from pathlib import Path
from typing import Any
import joblib,pandas as pd
from feature_engineering import MODEL_FEATURES,engineer_features
from utils import load_json

def predict_no_show(record: dict[str,Any],model_path: str|Path,threshold_path: str|Path,model_name="XGBoost"):
    model=joblib.load(model_path); threshold=float(load_json(threshold_path)[model_name]); featured=engineer_features(pd.DataFrame([record])); probability=float(model.predict_proba(featured[MODEL_FEATURES])[:,1][0]); return {"no_show_probability":probability,"threshold":threshold,"outreach_recommended":probability>=threshold,"model":model_name,"scope":"Decision support only; never deny care or penalize a patient."}
''')
write("src/generate_visuals.py", '''
"""Generate EDA and model visualizations."""
from __future__ import annotations
import argparse,json
from pathlib import Path
import joblib,matplotlib.pyplot as plt,numpy as np,pandas as pd
from sklearn.metrics import precision_recall_curve
from feature_engineering import MODEL_FEATURES,TARGET

def save(path):
    path.parent.mkdir(parents=True,exist_ok=True); plt.tight_layout(); plt.savefig(path.with_suffix(".png"),dpi=150,bbox_inches="tight"); plt.savefig(path.with_suffix(".svg"),bbox_inches="tight"); plt.close()

def main():
    p=argparse.ArgumentParser(); p.add_argument("--input",default="data/processed/model_features.csv"); a=p.parse_args(); root=Path(__file__).resolve().parents[1]; frame=pd.read_csv(a.input,parse_dates=["scheduled_datetime","appointment_date"]); v=root/"visuals"
    counts=frame[TARGET].value_counts().reindex([0,1]); plt.figure(figsize=(7,4.5)); plt.bar(["Attended","No-show"],counts.values); plt.title("Appointment Outcome Distribution"); plt.ylabel("Appointments"); save(v/"class_distribution")
    plt.figure(figsize=(7,4.5)); plt.hist(frame.age,bins=30); plt.title("Age Distribution"); plt.xlabel("Age"); save(v/"age_distribution")
    plt.figure(figsize=(7,4.5)); plt.hist(frame.waiting_days.clip(upper=60),bins=30); plt.title("Waiting Time Distribution"); plt.xlabel("Days"); save(v/"waiting_time_distribution")
    rates=frame.groupby("waiting_bucket",observed=False)[TARGET].mean(); plt.figure(figsize=(8,4.5)); plt.bar(rates.index.astype(str),rates.values); plt.xticks(rotation=35); plt.title("No-show Rate by Waiting-Time Bucket"); save(v/"no_show_by_waiting_bucket")
    rates=frame.groupby("sms_received")[TARGET].mean().reindex([0,1]); plt.figure(figsize=(6.5,4.5)); plt.bar(["No SMS","SMS Received"],rates.values); plt.title("Observed No-show Rate by SMS Status"); save(v/"no_show_by_sms")
    rates=frame.groupby("age_group",observed=False)[TARGET].mean(); plt.figure(figsize=(8,4.5)); plt.bar(rates.index.astype(str),rates.values); plt.xticks(rotation=30); plt.title("No-show Rate by Age Group"); save(v/"no_show_by_age_group")
    corr=frame[["age","waiting_days","schedule_hour","chronic_count","sms_received",TARGET]].corr(); plt.figure(figsize=(7,5.5)); plt.imshow(corr,cmap="coolwarm",vmin=-1,vmax=1); plt.colorbar(); plt.xticks(range(len(corr)),corr.columns,rotation=45,ha="right"); plt.yticks(range(len(corr)),corr.columns); plt.title("Numeric Feature Correlation"); save(v/"correlation_heatmap")
    metrics=json.loads((root/"reports/benchmark_metrics.json").read_text()); test=[r for r in metrics["models"] if r["split"]=="test"]; labels=[r["model"] for r in test]; x=np.arange(2); plt.figure(figsize=(8,4.8)); plt.bar(x-.18,[r["pr_auc"] for r in test],.36,label="PR-AUC"); plt.bar(x+.18,[r["recall"] for r in test],.36,label="Recall"); plt.xticks(x,labels); plt.legend(); plt.title("Model Comparison on Chronological Test Set"); save(v/"model_comparison")
    parts=joblib.load(root/"models/evaluation_partitions.joblib"); test_frame=parts["test"]; plt.figure(figsize=(7,5))
    for name,file in [("Logistic Regression","baseline_logistic_regression.joblib"),("XGBoost","advanced_xgboost.joblib")]:
        model=joblib.load(root/"models"/file); prob=model.predict_proba(test_frame[MODEL_FEATURES])[:,1]; precision,recall,_=precision_recall_curve(test_frame[TARGET],prob); plt.plot(recall,precision,label=name)
    plt.axhline(test_frame[TARGET].mean(),linestyle="--",label="Prevalence"); plt.xlabel("Recall"); plt.ylabel("Precision"); plt.title("Precision–Recall Curves"); plt.legend(); save(v/"precision_recall_curve")
    recommended=metrics["recommended_model"]; row=next(r for r in test if r["model"]==recommended); matrix=np.array(row["confusion_matrix"]); plt.figure(figsize=(5.5,5)); plt.imshow(matrix,cmap="Blues"); plt.xticks([0,1],["Attend","No-show"]); plt.yticks([0,1],["Attend","No-show"]); plt.xlabel("Predicted"); plt.ylabel("Actual");
    for i in range(2):
        for j in range(2): plt.text(j,i,f"{matrix[i,j]:,}",ha="center",va="center")
    plt.title(f"Confusion Matrix — {recommended}"); save(v/"confusion_matrix")
    advanced=joblib.load(root/"models/advanced_xgboost.joblib"); names=advanced.named_steps["preprocessor"].get_feature_names_out(); importance=advanced.named_steps["model"].feature_importances_; top=np.argsort(importance)[-15:]; plt.figure(figsize=(8,6)); plt.barh([names[i].replace("cat__","").replace("num__","").replace("bin__","") for i in top],importance[top]); plt.title("XGBoost Feature Importance"); save(v/"feature_importance")
if __name__=="__main__": main()
''')
write("run_pipeline.py", '''
"""Run the complete reproducible modeling pipeline."""
from __future__ import annotations
import argparse,subprocess,sys
from pathlib import Path

def run(command): print("$"," ".join(command)); subprocess.run(command,check=True)
def main():
    p=argparse.ArgumentParser(); p.add_argument("--input",default=""); p.add_argument("--demo-rows",type=int,default=30000); a=p.parse_args(); root=Path(__file__).resolve().parent; py=sys.executable; raw=Path(a.input) if a.input else root/"data/sample_data/demo_appointments.csv"
    if not a.input: run([py,str(root/"src/generate_demo_data.py"),"--rows",str(a.demo_rows),"--output",str(raw)])
    run([py,str(root/"src/preprocess.py"),"--input",str(raw),"--output",str(root/"data/processed/clean_appointments.csv"),"--quality-output",str(root/"reports/data_quality.json")]); run([py,str(root/"src/feature_engineering.py"),"--input",str(root/"data/processed/clean_appointments.csv"),"--output",str(root/"data/processed/model_features.csv")]); run([py,str(root/"src/train_model.py"),"--input",str(root/"data/processed/model_features.csv")]); run([py,str(root/"src/evaluate_model.py")]); run([py,str(root/"src/generate_visuals.py"),"--input",str(root/"data/processed/model_features.csv")])
if __name__=="__main__": main()
''')
write("app/app.py", '''
"""Streamlit no-show outreach decision-support app."""
from __future__ import annotations
import json,sys
from datetime import date,time
from pathlib import Path
import joblib,pandas as pd,streamlit as st
ROOT=Path(__file__).resolve().parents[1]; sys.path.insert(0,str(ROOT/"src")); from feature_engineering import MODEL_FEATURES,engineer_features
st.set_page_config(page_title="Medical Appointment No-Show Intelligence",page_icon="🏥",layout="wide")
@st.cache_resource
def load_models(): return joblib.load(ROOT/"models/baseline_logistic_regression.joblib"),joblib.load(ROOT/"models/advanced_xgboost.joblib"),json.loads((ROOT/"models/threshold_config.json").read_text())
baseline,advanced,thresholds=load_models(); metrics=json.loads((ROOT/"reports/benchmark_metrics.json").read_text()); page=st.sidebar.radio("Navigate",["Home","Project Overview","Prediction Interface","Model Performance","Visualizations","Business Insights","About"]); st.sidebar.warning("Decision-support demonstration only. Never deny or deprioritize care using this score.")
if page=="Home":
    st.title("Medical Appointment No-Show Intelligence"); c=st.columns(4); c[0].metric("Official appointments","110,527"); c[1].metric("Official columns","14"); c[2].metric("No-show prevalence","≈20.2%"); c[3].metric("Primary metric","PR-AUC"); st.info("Rerun on the official Kaggle CSV before quoting final performance.")
elif page=="Project Overview": st.header("Project Overview"); st.markdown("Supportive reminder and care-navigation prioritization using chronological validation. Patient and appointment identifiers are excluded from modeling.")
elif page=="Prediction Interface":
    st.header("Single-Appointment Outreach Score"); label=st.selectbox("Model",["XGBoost","Logistic Regression"]); model=advanced if label=="XGBoost" else baseline
    with st.form("score"):
        c1,c2,c3=st.columns(3)
        with c1: appt=st.date_input("Appointment date",date(2016,6,10)); wait=st.number_input("Waiting days",0,180,14); sched_time=st.time_input("Scheduling time",time(10,0)); age=st.number_input("Age",0,110,35)
        with c2: gender=st.selectbox("Gender code",["F","M"]); hood=st.selectbox("Neighbourhood",["JARDIM CAMBURI","MARIA ORTIZ","RESISTÊNCIA","JARDIM DA PENHA","CENTRO","OTHER"]); scholarship=st.selectbox("Scholarship",[0,1]); sms=st.selectbox("SMS received",[0,1])
        with c3: hypertension=st.selectbox("Hypertension",[0,1]); diabetes=st.selectbox("Diabetes",[0,1]); alcoholism=st.selectbox("Alcoholism",[0,1]); handicap=st.number_input("Handicap count",0,4,0)
        submitted=st.form_submit_button("Score appointment",use_container_width=True)
    if submitted:
        appointment=pd.Timestamp(appt,tz="UTC"); scheduled=appointment-pd.Timedelta(days=int(wait))+pd.Timedelta(hours=sched_time.hour,minutes=sched_time.minute); record={"gender":gender,"scheduled_datetime":scheduled,"appointment_date":appointment,"age":age,"neighbourhood":hood,"scholarship":scholarship,"hypertension":hypertension,"diabetes":diabetes,"alcoholism":alcoholism,"handicap":handicap,"sms_received":sms,"waiting_days":wait}; featured=engineer_features(pd.DataFrame([record])); probability=float(model.predict_proba(featured[MODEL_FEATURES])[:,1][0]); threshold=float(thresholds[label]); st.metric("Estimated no-show probability",f"{probability:.1%}"); st.metric("Outreach threshold",f"{threshold:.1%}"); st.warning("Suggested workflow: supportive reminder or assistance—not denial of care.") if probability>=threshold else st.success("Suggested workflow: standard reminder process.")
elif page=="Model Performance":
    st.header("Model Performance"); st.dataframe(pd.DataFrame([{k:v for k,v in r.items() if k!="confusion_matrix"} for r in metrics["models"]]),hide_index=True,use_container_width=True); st.markdown(f"**Validation-selected demonstration model:** {metrics['recommended_model']}"); st.caption(metrics["benchmark_scope"])
    for name in ["model_comparison.svg","precision_recall_curve.svg","confusion_matrix.svg"]: st.image(str(ROOT/"visuals"/name),use_container_width=True)
elif page=="Visualizations":
    st.header("Exploratory Visualizations")
    for name in ["class_distribution.svg","age_distribution.svg","waiting_time_distribution.svg","no_show_by_waiting_bucket.svg","no_show_by_sms.svg","no_show_by_age_group.svg","correlation_heatmap.svg","feature_importance.svg"]: st.image(str(ROOT/"visuals"/name),caption=name.replace("_"," ").replace(".svg","").title(),use_container_width=True)
elif page=="Business Insights": st.header("Business Insights"); st.markdown("1. Use risk bands for supportive outreach.\n2. Choose thresholds from real capacity and costs.\n3. Treat SMS as observational, not causal.\n4. Monitor subgroup performance and patient experience.\n5. Never restrict access using a score.")
else: st.header("About"); st.markdown("Built by **Rachel Oyeyemi** as a recruiter-ready healthcare machine-learning portfolio project.")
''')
write("data/raw/README.md", "Run `python src/download_data.py --source auto` after authenticating to Kaggle. Raw healthcare data is excluded from version control.\n")
write("data/processed/README.md", "Processed CSV files are generated reproducibly and excluded from version control.\n")
write("models/README.md", "Committed models are deterministic demonstration artifacts. Retrain on the official CSV before quoting results.\n")
write("PROJECT_CHARTER.md", '''
# Project Charter — Medical Appointment No Shows

## Business Problem
Estimate appointment-level no-show risk early enough to prioritize considerate reminders, transportation support, rescheduling, or human care navigation.

## Objectives
Build a leakage-aware classification pipeline; compare Logistic Regression and XGBoost; select transparent outreach thresholds; deliver Streamlit, reports, tests, and presentation materials.

## Stakeholders
Clinic operations, schedulers, care-navigation teams, data teams, privacy/compliance leaders, and patients.

## Success Metrics
PR-AUC, recall, precision, F1, ROC-AUC, balanced accuracy, MCC, outreach volume, appointments recovered, and subgroup equity.

## Architecture
`Kaggle → validation → cleaning → features → chronological partitions → models → threshold policy → evaluation → Streamlit`

## Responsible Use
Never deny, delay, or deprioritize care based on a model score. Require privacy review, subgroup testing, human oversight, and intervention evaluation.
''')
write("EDA_REPORT.md", '''
# Exploratory Data Analysis Report

The widely reproduced raw profile contains 110,527 Brazilian appointments and 14 columns. `No-show = Yes` is the positive class and represents roughly 20.2% of records. Published reproductions report no missing values and no exact duplicate rows.

## Quality Issues
- At least one invalid negative age
- `Handcap` values above one
- Spelling inconsistencies in source column names
- Potential negative waiting times after date parsing
- Identifier and repeat-patient leakage risk
- SMS status is observational and may be confounded by lead time

## Analysis
The project examines class balance, age, waiting days, SMS status, age groups, neighbourhoods, chronic-condition indicators, correlations, and outliers. Accuracy is insufficient because attendance is the majority outcome.
''')
write("MODEL_COMPARISON.md", '''
# Model Comparison

## Baseline
Class-weighted Logistic Regression with standardized numeric variables and one-hot categories. Transparent, fast, and explainable.

## Advanced
Regularized XGBoost with positive-class weighting. Captures nonlinear thresholds and interactions.

## Demonstration Results
| Model | Test PR-AUC | Precision | Recall | F1 | ROC-AUC |
|---|---:|---:|---:|---:|---:|
| Logistic Regression | 0.468 | 0.294 | 0.861 | 0.438 | 0.765 |
| XGBoost | 0.454 | 0.315 | 0.804 | 0.453 | 0.758 |

XGBoost wins validation PR-AUC by a very small margin; Logistic Regression ranks better and has higher recall on the synthetic test set. Preserve both as champion/challenger candidates and rerun on official data.
''')
write("MODEL_EVALUATION.md", '''
# Model Evaluation

Appointments are split chronologically into 70% training, 15% validation, and 15% test partitions. Thresholds are selected only on validation data using illustrative costs of 5 for a missed no-show and 1 for unnecessary outreach.

Primary metric: PR-AUC. Supporting metrics: accuracy, precision, recall, F1, ROC-AUC, balanced accuracy, MCC, specificity, outreach rate, and confusion matrix.

The demonstration threshold is intentionally recall-oriented. A real clinic must choose thresholds from actual capacity, costs, patient preferences, and equity analysis.
''')
write("BUSINESS_RECOMMENDATIONS.md", '''
# Business Recommendations

- Use low, medium, and high support bands rather than punitive decisions.
- Offer reminders, rescheduling, transportation, language, childcare, or telehealth support where available.
- Track appointments recovered per 1,000 contacts, contact fatigue, opt-outs, and subgroup outcomes.
- Run randomized holdouts to measure incremental intervention impact.
- Never deny, delay, or deprioritize care using the score.

## Risks
Privacy, socioeconomic proxy bias, feedback loops, calibration drift, and inequitable access.

## Future Opportunities
Time-safe attendance history, appointment type, provider, travel distance, causal reminder analysis, calibration, capacity-aware optimization, and fairness-constrained thresholds.
''')
write("DATA_DICTIONARY.md", '''
# Data Dictionary

| Source | Clean name | Use |
|---|---|---|
| PatientId | patient_id | Excluded identifier |
| AppointmentID | appointment_id | Excluded identifier |
| Gender | gender | Categorical |
| ScheduledDay | scheduled_datetime | Calendar features |
| AppointmentDay | appointment_date | Calendar features and split |
| Age | age | Numeric and age group |
| Neighbourhood | neighbourhood | Categorical; governance review |
| Scholarship | scholarship | Binary |
| Hipertension | hypertension | Binary |
| Diabetes | diabetes | Binary |
| Alcoholism | alcoholism | Binary |
| Handcap | handicap | Recorded count |
| SMS_received | sms_received | Binary, observational |
| No-show | no_show | Target: Yes = 1 |
''')
write("docs/TECHNICAL_ARCHITECTURE.md", "# Technical Architecture\n\nIngestion → validation → preprocessing → feature engineering → chronological training/validation/test → threshold selection → evaluation → Streamlit. Raw and processed patient-level CSVs are excluded from Git.\n")
write("docs/MODEL_CARD.md", '''
# Model Card

## Intended Use
Educational outreach-prioritization demonstration.

## Prohibited Use
Denying care, penalizing patients, clinical diagnosis, or autonomous decision-making.

## Limitations
Public historical data, short time coverage, geographic specificity, limited operational features, observational SMS status, and no validated intervention outcomes.

## Governance
Privacy review, calibration, subgroup evaluation, human override, patient-centered policy, drift monitoring, and documented threshold approval are required.
''')
write("RESUME_LINKEDIN_INTERVIEW.md", '''
# Resume, LinkedIn, and Interview Materials

## Resume Bullets
- Built an end-to-end medical appointment no-show pipeline with schema validation, leakage-aware features, chronological evaluation, and automated tests.
- Compared class-weighted Logistic Regression and XGBoost using PR-AUC, recall, precision, F1, ROC-AUC, balanced accuracy, and cost-aware thresholds.
- Developed a responsible Streamlit decision-support app, executive presentation, model card, CI workflow, and patient-access guardrails.

## LinkedIn Description
Built a complete Medical Appointment No-Show prediction portfolio focused on supportive outreach rather than patient blame. The project includes production-style engineering, temporal validation, Logistic Regression and XGBoost, threshold economics, automated tests, Streamlit, business recommendations, and responsible-AI documentation.

## Interview Talking Points
Class imbalance; chronological validation; identifier leakage; PR-AUC versus ROC-AUC; threshold selection; observational SMS status; neighbourhood fairness; champion/challenger models.

## Sample Questions
**Why PR-AUC?** It focuses on the less common no-show class and the precision/recall trade-off.  
**How did you prevent leakage?** I excluded identifiers and used only pre-appointment features with chronological partitions.  
**How would this be used safely?** Only for optional reminders and support, with human oversight and no access restrictions.
''')
write("README.md", '''
# Medical Appointment No Shows

[![CI](https://github.com/Rachel-Oyeyemi/Medical-Appointment-No-Shows/actions/workflows/ci.yml/badge.svg)](https://github.com/Rachel-Oyeyemi/Medical-Appointment-No-Shows/actions)

A recruiter-ready machine-learning portfolio project for appointment no-show outreach prioritization using data engineering, chronological validation, Logistic Regression, XGBoost, cost-sensitive thresholds, Streamlit, tests, executive reporting, and responsible-AI controls.

> Committed metrics use deterministic synthetic appointments so the repository runs immediately. Download the official Kaggle CSV and rerun before quoting final performance.

## Dataset
Kaggle: `joniarroba/noshowappointments`. Widely reproduced profile: 110,527 Brazilian appointments, 14 columns, target `No-show`, and approximately 20.2% no-shows. Published reproductions report no missing values and no exact duplicate rows. Quality checks include an invalid negative age, non-binary `Handcap` values, date parsing, and possible negative waiting times.

## Methodology
1. Download and validate the CSV.
2. Standardize names, dates, target, and data types.
3. Engineer waiting time, calendar, age-group, waiting-bucket, and chronic-condition features.
4. Exclude PatientId and AppointmentID.
5. Split chronologically 70/15/15.
6. Train class-weighted Logistic Regression and XGBoost.
7. Select validation thresholds from transparent outreach costs.
8. Evaluate PR-AUC, recall, precision, F1, ROC-AUC, balanced accuracy, MCC, specificity, and outreach volume.

## Demonstration Results
| Model | Test PR-AUC | Precision | Recall | F1 | ROC-AUC | Outreach Rate |
|---|---:|---:|---:|---:|---:|---:|
| Logistic Regression | 0.468 | 0.294 | 0.861 | 0.438 | 0.765 | 58.9% |
| XGBoost | 0.454 | 0.315 | 0.804 | 0.453 | 0.758 | 51.3% |

XGBoost is the validation-selected demo model by a very small margin. Mixed test results are retained honestly.

## Run
```bash
git clone https://github.com/Rachel-Oyeyemi/Medical-Appointment-No-Shows.git
cd Medical-Appointment-No-Shows
python -m venv .venv
# Windows: .venv\\Scripts\\activate
# macOS/Linux: source .venv/bin/activate
pip install -r requirements.txt
python run_pipeline.py
streamlit run app/app.py
pytest -q
```

Official data: `python src/download_data.py --source auto` then `python run_pipeline.py --input data/raw/appointments.csv`.

## Structure
`app/`, `data/`, `docs/`, `models/`, `notebooks/`, `presentation/`, `reports/`, `src/`, `tests/`, and `visuals/`, plus the requested charter, EDA, evaluation, recommendations, and career materials.

## Responsible Use
Never deny, delay, or deprioritize care using a model score. Production requires institutional validation, privacy and security review, calibration, subgroup testing, intervention experiments, monitoring, and human override.
''')
write("presentation/SLIDE_NOTES.md", "# Speaker Notes\n\nProblem → dataset → EDA → modeling → results → insights → recommendations → future work → responsible conclusion. Emphasize supportive outreach and synthetic benchmark scope.\n")
write("presentation/create_presentation.py", '''
"""Generate the 10-slide executive PowerPoint."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches,Pt
ROOT=Path(__file__).resolve().parents[1]; NAVY=RGBColor(18,35,63); DARK=RGBColor(38,45,55); MUTED=RGBColor(92,104,120)
def title(slide,text,sub=""):
    box=slide.shapes.add_textbox(Inches(.7),Inches(.4),Inches(12),Inches(.72)); p=box.text_frame.paragraphs[0]; p.text=text; p.font.size=Pt(28); p.font.bold=True; p.font.color.rgb=NAVY
    if sub: box=slide.shapes.add_textbox(Inches(.72),Inches(1.06),Inches(11.8),Inches(.42)); p=box.text_frame.paragraphs[0]; p.text=sub; p.font.size=Pt(11); p.font.color.rgb=MUTED
def bullets(slide,items,x=.8,y=1.55,w=5.65,h=4.9,size=18):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear()
    for i,item in enumerate(items): p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=item; p.font.size=Pt(size); p.font.color.rgb=DARK; p.space_after=Pt(10)
def chart(slide,name):
    path=ROOT/"visuals"/name
    if path.exists(): slide.shapes.add_picture(str(path),Inches(6.65),Inches(1.5),width=Inches(5.85),height=Inches(4.9))
def main():
    metrics=json.loads((ROOT/"reports/benchmark_metrics.json").read_text()); test={r["model"]:r for r in metrics["models"] if r["split"]=="test"}; prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
    slide=prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=NAVY; box=slide.shapes.add_textbox(Inches(.8),Inches(1.15),Inches(11.7),Inches(1.2)); p=box.text_frame.paragraphs[0]; p.text="Medical Appointment No Shows"; p.font.size=Pt(39); p.font.bold=True; p.font.color.rgb=RGBColor(255,255,255); box=slide.shapes.add_textbox(Inches(.82),Inches(2.55),Inches(11),Inches(1)); p=box.text_frame.paragraphs[0]; p.text="Outreach prioritization, threshold economics, and responsible healthcare AI"; p.font.size=Pt(21); p.font.color.rgb=RGBColor(210,222,241)
    specs=[("2. Business Problem","Support attendance without restricting access",["Missed appointments waste capacity and can delay care.","Blanket outreach consumes staff time.","Prioritize supportive reminders and care navigation."],"class_distribution.png"),("3. Dataset","Kaggle Medical Appointment No Shows",["110,527 appointment records","14 source columns","Approximately 20.2% no-shows","Brazilian context and limited time coverage"],"age_distribution.png"),("4. Exploratory Analysis","Lead time, age, location, and policy context",["Longer waits are associated with risk.","Class imbalance makes accuracy misleading.","SMS status is observational.","Identifiers and invalid records require controls."],"no_show_by_waiting_bucket.png"),("5. Modeling","Transparent baseline versus nonlinear challenger",["Class-weighted Logistic Regression","Regularized XGBoost","Chronological 70/15/15 split","Validation-only threshold selection"],"feature_importance.png"),("6. Results","Deterministic demonstration benchmark",[f"Logistic test PR-AUC: {test['Logistic Regression']['pr_auc']:.3f}",f"XGBoost test PR-AUC: {test['XGBoost']['pr_auc']:.3f}",f"Validation-selected model: {metrics['recommended_model']}","Official Kaggle rerun required"],"model_comparison.png"),("7. Key Insights","Ranking and operating policy are separate",["High recall can require substantial outreach.","XGBoost reduces false alerts at its threshold.","Logistic ranks better on the demo test set.","Keep champion and challenger models."],"precision_recall_curve.png"),("8. Recommendations","Use patient-centered risk bands",["Low risk: standard reminders","Medium risk: automated options","High risk: human support or rescheduling","Never deny care based on a score"],None),("9. Future Work","Move to institutional decision support",["Time-safe attendance history","Appointment, provider, travel, and language features","Calibration and temporal cross-validation","Randomized intervention evaluation"],None)]
    for heading,sub,items,img in specs:
        slide=prs.slides.add_slide(blank); title(slide,heading,sub); bullets(slide,items) if img else bullets(slide,items,1,1.7,11.2,4.7,21); chart(slide,img) if img else None
    slide=prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=NAVY; box=slide.shapes.add_textbox(Inches(.8),Inches(1.15),Inches(11.7),Inches(.9)); p=box.text_frame.paragraphs[0]; p.text="10. Conclusion"; p.font.size=Pt(34); p.font.bold=True; p.font.color.rgb=RGBColor(255,255,255); box=slide.shapes.add_textbox(Inches(1),Inches(2.35),Inches(11.2),Inches(2.6)); p=box.text_frame.paragraphs[0]; p.text="Responsible no-show prediction combines leakage-aware modeling, transparent outreach thresholds, patient-centered interventions, equity monitoring, and human oversight."; p.font.size=Pt(27); p.font.color.rgb=RGBColor(220,230,245); p.alignment=PP_ALIGN.CENTER; prs.save(Path(__file__).with_name("Medical_Appointment_No_Shows_Executive_Presentation.pptx"))
if __name__=="__main__": main()
''')
write("tests/test_pipeline.py", '''
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"src"))
from generate_demo_data import generate_demo
from preprocess import preprocess_dataframe
from feature_engineering import MODEL_FEATURES,engineer_features
from train_model import build_baseline
from evaluate_model import choose_threshold

def test_clean_and_features():
    raw=generate_demo(300); raw.loc[0,"Age"]=-1; clean,report=preprocess_dataframe(raw); featured=engineer_features(clean); assert report["invalid_age_rows_removed"]==1; assert set(MODEL_FEATURES).issubset(featured.columns); assert not featured[MODEL_FEATURES].isna().any().any()
def test_model_probabilities():
    clean,_=preprocess_dataframe(generate_demo(1200)); features=engineer_features(clean); model=build_baseline().fit(features[MODEL_FEATURES],features.no_show); probabilities=model.predict_proba(features[MODEL_FEATURES].head(20))[:,1]; assert ((probabilities>=0)&(probabilities<=1)).all()
def test_threshold():
    import numpy as np
    threshold,rows=choose_threshold(np.array([0,0,0,1,1,1]),np.array([.1,.2,.4,.45,.7,.9])); assert .05<=threshold<=.90 and len(rows)>10
''')
write(".github/workflows/ci.yml", '''
name: CI
on: [push, pull_request]
jobs:
  test:
    runs-on: ubuntu-latest
    steps:
      - uses: actions/checkout@v4
      - uses: actions/setup-python@v5
        with: {python-version: "3.11"}
      - run: pip install -r requirements.txt
      - run: pytest -q
''')

# Create notebooks.
notebook_cells = {
    "01_data_exploration.ipynb": ["# 01 — Data Exploration", "from pathlib import Path\nimport sys,pandas as pd\nsys.path.append('../src')\nfrom generate_demo_data import generate_demo\nfrom preprocess import preprocess_dataframe\npath=Path('../data/raw/appointments.csv')\nraw=pd.read_csv(path) if path.exists() else generate_demo(5000)\nclean,quality=preprocess_dataframe(raw)\nquality", "clean[['age','waiting_days','sms_received','no_show']].describe()"],
    "02_preprocessing.ipynb": ["# 02 — Preprocessing", "import sys\nsys.path.append('../src')\nfrom generate_demo_data import generate_demo\nfrom preprocess import preprocess_dataframe\nfrom feature_engineering import engineer_features\nclean,quality=preprocess_dataframe(generate_demo(5000))\nfeatures=engineer_features(clean)\nquality,features.head()"],
    "03_modeling.ipynb": ["# 03 — Modeling", "import sys\nsys.path.append('../src')\nfrom generate_demo_data import generate_demo\nfrom preprocess import preprocess_dataframe\nfrom feature_engineering import engineer_features\nfrom train_model import train_models\nclean,_=preprocess_dataframe(generate_demo(8000))\nfeatures=engineer_features(clean)\nbaseline,advanced,train,validation,test,weight=train_models(features)\nlen(train),len(validation),len(test)"],
    "04_business_insights.ipynb": ["# 04 — Business Insights", "import sys\nsys.path.append('../src')\nfrom generate_demo_data import generate_demo\nfrom preprocess import preprocess_dataframe\nfrom feature_engineering import engineer_features\nclean,_=preprocess_dataframe(generate_demo(10000))\nfeatures=engineer_features(clean)\nfeatures.groupby('waiting_bucket',observed=False)['no_show'].agg(['count','mean'])", "print('Use risk bands for support. Never deny or deprioritize care from a score.')"],
}
for name, cells in notebook_cells.items():
    nb=nbf.v4.new_notebook(); nb.metadata.kernelspec={"display_name":"Python 3","language":"python","name":"python3"}
    nb.cells=[nbf.v4.new_markdown_cell(cells[0])]+[nbf.v4.new_code_cell(cell) for cell in cells[1:]]
    nbf.write(nb,ROOT/"notebooks"/name)

print("Project source files generated")
